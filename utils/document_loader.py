import io
import pandas as pd

def extract_text(uploaded_file):
    name = uploaded_file.name.lower()
    raw = uploaded_file.getvalue()
    if name.endswith((".txt", ".md")):
        return raw.decode("utf-8", errors="replace")
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if name.endswith(".csv"):
        return pd.read_csv(io.BytesIO(raw)).to_csv(index=False)
    if name.endswith(".xlsx"):
        book = pd.ExcelFile(io.BytesIO(raw))
        parts = []
        for sheet in book.sheet_names:
            df = pd.read_excel(book, sheet_name=sheet)
            parts.append(f"Sheet: {sheet}\n{df.to_csv(index=False)}")
        return "\n\n".join(parts)
    raise ValueError("Unsupported file type")
